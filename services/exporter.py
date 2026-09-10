from datetime import datetime
from pathlib import Path
from typing import List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from core.config import settings
from core.logger import logger
from services.posnext.models import ParsedShow


class ExcelExporter:
    """Генератор профессионально оформленных Excel отчетов (.xlsx)."""

    def __init__(self, output_dir: str = settings.REPORTS_DIR):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(self, shows: List[ParsedShow], filename_prefix: str = "pos_solds_report") -> Path:
        """
        Создает .xlsx файл с таблицей:
        Show Name | Venue | Date & Time | Sold | Active On Hand
        """
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        filename = f"{filename_prefix}_{timestamp}.xlsx"
        filepath = self.output_dir / filename

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "POSNext Solds"

        # Включаем сетку
        ws.views.sheetView[0].showGridLines = True

        # Стили оформления
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        data_font = Font(name="Calibri", size=11)
        zebra_fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        align_left = Alignment(horizontal="left", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        # Заголовки таблицы строго по требованиям пользователя
        headers = ["Show Name", "Venue", "Date & Time", "Sold", "Active On Hand"]
        ws.append(headers)

        # Стилизуем шапку
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Заполняем строки данных
        for row_idx, show in enumerate(shows, start=2):
            ws.row_dimensions[row_idx].height = 22
            current_fill = zebra_fill if row_idx % 2 == 0 else white_fill

            # Формируем красивое название площадки (с городом/штатом если есть)
            venue_display = f"{show.venue} ({show.city_state})" if show.city_state else show.venue

            row_data = [
                (show.event_name, align_left, "@"),
                (venue_display, align_left, "@"),
                (show.event_datetime_str, align_center, "@"),
                (show.sold_count, align_center, "#,##0"),
                (show.active_on_hand, align_center, "#,##0"),
            ]

            for col_idx, (val, alignment, num_format) in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.fill = current_fill
                cell.alignment = alignment
                cell.border = thin_border
                cell.number_format = num_format

        # Автоматическая настройка ширины колонок с запасом
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        # Включаем автофильтр для всей таблицы
        if shows:
            last_col_letter = get_column_letter(len(headers))
            ws.auto_filter.ref = f"A1:{last_col_letter}{len(shows) + 1}"

        wb.save(filepath)
        logger.info(f"Excel отчет успешно сохранен: {filepath}")
        return filepath
